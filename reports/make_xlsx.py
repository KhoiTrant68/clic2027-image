"""CLIC 2027 proposal as a simple-language Excel report."""
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference, ScatterChart, Series
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = sys.argv[1]
F = "Arial"
TITLE = Font(name=F, size=14, bold=True)
SUB = Font(name=F, size=10, italic=True, color="666666")
HEAD = Font(name=F, size=10, bold=True, color="FFFFFF")
BODY = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
NOTE_FILL = PatternFill("solid", fgColor="F2F2F2")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
GOOD = PatternFill("solid", fgColor="E2EFDA")
WARN = PatternFill("solid", fgColor="FFF2CC")
BAD = PatternFill("solid", fgColor="F8CBAD")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")

wb = Workbook()


def sheet(name, title, subtitle, widths):
    ws = wb.active if wb.active.title == "Sheet" else wb.create_sheet()
    ws.title = name
    ws["A1"], ws["A1"].font = title, TITLE
    ws["A2"], ws["A2"].font = subtitle, SUB
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False
    return ws


def table(ws, row, headers, rows, fmt=None):
    """Write a header row + body rows starting at `row`; return the next free row."""
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font, cell.fill, cell.border, cell.alignment = HEAD, HEAD_FILL, BOX, WRAP
    for r, vals in enumerate(rows, row + 1):
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font, cell.border, cell.alignment = BODY, BOX, WRAP
            if fmt and c in fmt:
                cell.number_format = fmt[c]
    ws.freeze_panes = None
    return row + len(rows) + 2


def note(ws, row, text, ncols):
    ws.cell(row=row, column=1, value=text).font = SUB
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    ws.cell(row=row, column=1).alignment = WRAP
    ws.row_dimensions[row].height = 30
    return row + 2


def section(ws, row, text):
    ws.cell(row=row, column=1, value=text).font = Font(name=F, size=11, bold=True, color="1F4E79")
    return row + 1


# 1. Summary ---------------------------------------------------------------------
ws = sheet("Tóm tắt", "Đề xuất cách làm cho CLIC 2027 (ảnh, GPU track)", "Khoi Tran Minh · cập nhật 03/10/2026", [28, 90])
rows = [
    ("Mục tiêu", "Đứng hạng 1 track ảnh GPU của CLIC 2027. Hạn nộp decoder: 01/03/2027."),
    ("Ý chính", "Dùng 2 cách nén và chọn cho từng ảnh: (1) ảnh thường: tạo ảnh bằng AI rồi sửa lại bằng một lớp bù sai số; (2) ảnh màn hình: dùng codec VVC-SCC có sẵn."),
    ("Vì sao không dùng AI một mình", "Bộ mã hóa ảnh DC-AE (của SANA), dù không nén gì, chỉ đạt 22.86 dB. Đội thắng năm 2025 đạt 23.99 dB. Nên cần thêm lớp bù sai số."),
    ("Vì sao vẫn dùng AI", "Ban giám khảo là người thật, chọn ảnh nhìn đẹp hơn. Năm 2025, đội có PSNR thấp lại thắng. Ảnh tạo bằng AI có LPIPS tốt hơn VTM khoảng 4 lần."),
    ("3 ràng buộc lớn", "Decoder bị khóa trước vòng test · máy chủ phải giải mã ra đúng từng bit · 25% bài chậm nhất bị loại."),
    ("Đã làm xong", "Đo trần chất lượng, so với VTM/VTM-SCC, kiểm tra tính khớp bit, viết lại DC-AE/SANA không cần diffusers, codec S1 chạy được."),
    ("Rủi ro lớn nhất", "Decoder đang chậm hơn mục tiêu khoảng 9–13 lần. Chất lượng thật của phần AI ở bitrate thấp còn kém hơn dự tính."),
    ("Cần anh/chị góp ý", "Xem sheet \"Câu hỏi\". Có cột trống màu vàng để ghi ý kiến."),
]
r = table(ws, 4, ["Mục", "Nội dung"], rows)
for i in range(5, 5 + len(rows)):
    ws.cell(row=i, column=1).font = BOLD
r = section(ws, r, "Các sheet trong file")
table(ws, r, ["Sheet", "Có gì"], [
    ("Luật thi", "Luật của cuộc thi và điều đó ảnh hưởng thế nào tới cách làm"),
    ("Kết quả", "Các con số đã đo được, kèm biểu đồ"),
    ("Cách làm", "Các phần của hệ thống, mỗi phần làm gì"),
    ("Lộ trình", "Việc gì làm khi nào"),
    ("Rủi ro", "Điều có thể sai và cách xử lý"),
    ("Câu hỏi", "Những điểm cần anh/chị góp ý"),
])

# 2. Rules ------------------------------------------------------------------------
ws = sheet("Luật thi", "Luật thi và ảnh hưởng tới cách làm", "Nguồn: clic2027.compression.cc và devkit, đọc ngày 29/09/2026", [22, 46, 52])
r = table(ws, 4, ["Luật", "Giá trị", "Ảnh hưởng"], [
    ("Mức nén", "0.075 / 0.15 / 0.3 bpp, phải nộp đủ cả 3", "Một mô hình phải chạy được ở cả 3 mức"),
    ("Ngân sách", "Tính trên cả 30 ảnh (xem bảng dưới)", "Được chuyển bit từ ảnh dễ sang ảnh khó"),
    ("Cách chấm", "Người thật so từng cặp ảnh, xếp hạng Elo", "Ảnh phải đẹp, nhưng không được bịa nội dung"),
    ("Máy chấm", "GPU L4 24 GB, 2 CPU, RAM 12 GB, torch 2.6, không có diffusers", "Decoder viết bằng torch thuần"),
    ("Tốc độ", "25% bài giải mã chậm nhất bị loại", "Mục tiêu: 30 ảnh trong 20 giây"),
    ("Dung lượng decoder", "Tối đa 4 GB", "Dùng SANA-0.6B (khoảng 1.6 GB)"),
    ("Khóa decoder", "Decoder vòng test phải giống hệt bản nộp trước 01/03", "Chỉ encoder còn sửa được sau 01/03"),
    ("Số phiên bản", "Tối đa 3, khai báo từ vòng validation", "Làm 3 bản: thiên về giống gốc, thiên về đẹp, kết hợp"),
])
r = section(ws, r, "Ngân sách byte cho 30 ảnh validation")
ws.cell(row=r, column=1, value="Tổng số pixel").font = BOLD
ws.cell(row=r, column=2, value=86245376).number_format = "#,##0"
ws.cell(row=r, column=2).font = Font(name=F, size=10, color="0000FF")
ws.cell(row=r, column=3, value="Đo từ 30 ảnh validation (results/clic_b/b_results)").font = SUB
px = r
r = table(ws, r + 2, ["Mức nén (bpp)", "Ngân sách (byte)", "Cách tính"], [
    (0.075, f"=ROUNDDOWN(A{r + 3}*$B${px}/8,0)", "bpp × số pixel ÷ 8"),
    (0.15, f"=ROUNDDOWN(A{r + 4}*$B${px}/8,0)", "bpp × số pixel ÷ 8"),
    (0.3, f"=ROUNDDOWN(A{r + 5}*$B${px}/8,0)", "bpp × số pixel ÷ 8"),
], fmt={1: "0.000", 2: "#,##0"})
for i in (r - 4, r - 3, r - 2):
    ws.cell(row=i, column=1).font = Font(name=F, size=10, color="0000FF")
r = section(ws, r, "Mốc cần vượt (năm 2025, mức 0.075 bpp, cùng bộ ảnh)")
table(ws, r, ["Đội", "Elo", "PSNR (dB)", "Ghi chú"], [
    ("Vcoder (hạng 1)", 1929, 23.99, "Decoder 1.07 GB, 24 giây cho 30 ảnh"),
    ("Evolve", 1903, None, ""),
    ("VTM (codec chuẩn)", 1405, 26.52, "PSNR cao nhất nhưng Elo thấp nhất"),
], fmt={3: "0.00"})

# 3. Results ----------------------------------------------------------------------
ws = sheet("Kết quả", "Kết quả đã chạy (29/09 – 03/10/2026)", "Tất cả chạy trên Kaggle. Số màu xanh là số đo được; số màu đen là công thức.", [30, 14, 14, 14, 14, 40])
blue = Font(name=F, size=10, color="0000FF")

r = section(ws, 4, "1. Chất lượng tối đa (trần) trên 30 ảnh validation")
r0 = r
ceil = [("VTM @0.075 bpp", 26.52), ("HM @0.075 bpp", 25.83), ("SD-VAE f8 (trần)", 24.73),
        ("x265 @0.075 bpp", 24.58), ("Vcoder @0.075 bpp (hạng 1)", 23.99), ("DC-AE f32 (trần)", 22.86)]
rows = [(n, v, f"=B{r0 + 1 + i}-$B${r0 + 5}") for i, (n, v) in enumerate(ceil)]
r = table(ws, r0, ["Cấu hình", "PSNR (dB)", "So với Vcoder (dB)"], rows, fmt={2: "0.00", 3: "+0.00;-0.00;0.00"})
for i in range(r0 + 1, r0 + 1 + len(ceil)):
    ws.cell(row=i, column=2).font = blue
r = note(ws, r - 1, "Trần = mã hóa rồi giải mã, không nén. DC-AE thấp hơn Vcoder ~1.1 dB ngay cả khi không nén → cần thêm lớp bù sai số.", 6)

r = section(ws, r, "2. So sánh 3 cách nén (7 ảnh đại diện)")
b6 = [("Màn hình", 0.075, 32.93, 34.00, 30.98, 0.1257, 0.1161, 0.0454),
      ("Màn hình", 0.15, 36.30, 38.38, 34.25, 0.0794, 0.0656, 0.0343),
      ("Màn hình", 0.3, 39.84, 43.17, 38.03, 0.0455, 0.0417, 0.0231),
      ("Tự nhiên", 0.075, 24.82, 24.97, 24.61, 0.4009, 0.3910, 0.1048),
      ("Tự nhiên", 0.15, 27.44, 27.49, 26.07, 0.2636, 0.2669, 0.0923),
      ("Tự nhiên", 0.3, 29.97, 30.57, 28.29, 0.1696, 0.1610, 0.0684)]
for c in range(7, 10):
    ws.column_dimensions[get_column_letter(c)].width = 14
ws.column_dimensions["J"].width = 18
hb = r
rows = [(g, b, p1, p2, p3, l1, l2, l3, f"=F{hb + 1 + i}/H{hb + 1 + i}") for i, (g, b, p1, p2, p3, l1, l2, l3) in enumerate(b6)]
r = table(ws, hb, ["Nhóm ảnh", "bpp", "PSNR VTM", "PSNR VTM-SCC", "PSNR AI + bù", "LPIPS VTM", "LPIPS VTM-SCC", "LPIPS AI + bù", "LPIPS: VTM kém hơn (lần)"],
          rows, fmt={2: "0.000", 3: "0.00", 4: "0.00", 5: "0.00", 6: "0.000", 7: "0.000", 8: "0.000", 9: "0.0\"×\""})
for i in range(hb + 1, hb + 1 + len(b6)):
    for c in range(3, 9):
        ws.cell(row=i, column=c).font = blue
r = note(ws, r - 1, "PSNR cao hơn là tốt; LPIPS thấp hơn là tốt. \"AI + bù\" giả định phần AI tốn 0.03 bpp và đạt trần DC-AE (lạc quan), phần bù mã hóa bằng VTM (bi quan). Kết luận: ảnh màn hình → VVC-SCC; ảnh tự nhiên → AI + bù.", 9)

r = section(ws, r, "3. Codec S1 trên Kodak (24 ảnh), so với trần DC-AE 23.66 dB")
ws.cell(row=r, column=1, value="Trần DC-AE (dB)").font = BOLD
ws.cell(row=r, column=2, value=23.66).font = blue
ceil_row = r
r += 2
s1 = [("r0", 0.0155, 17.83, 0.5229), ("r1", 0.0204, 18.57, 0.4485), ("r2", 0.0342, 19.62, 0.3091), ("r3", 0.0531, 20.71, 0.2058),
      ("r4", 0.0710, 21.70, 0.1521), ("r5", 0.0885, 22.45, 0.1216), ("r6", 0.1057, 22.99, 0.1051), ("r7", 0.1215, 23.29, 0.0980)]
hs = r
rows = [(n, b, p, l, f"=C{hs + 1 + i}-$B${ceil_row}") for i, (n, b, p, l) in enumerate(s1)]
r = table(ws, hs, ["Mức rate", "bpp", "PSNR (dB)", "LPIPS", "Cách trần (dB)"], rows, fmt={2: "0.0000", 3: "0.00", 4: "0.000", 5: "+0.00;-0.00;0.00"})
for i in range(hs + 1, hs + 1 + len(s1)):
    for c in (2, 3, 4):
        ws.cell(row=i, column=c).font = blue
ws.cell(row=hs + 3, column=1).fill = WARN
r = note(ws, r - 1, "S1 được train bằng MSE nên ảnh còn mờ; bước AI sau (bridge) mới làm ảnh đẹp. Ở 0.034 bpp (dòng tô vàng) còn cách trần 4 dB → số \"AI + bù\" ở mục 2 đang lạc quan.", 6)

ch = ScatterChart()
ch.title, ch.style = "S1 trên Kodak: PSNR theo bpp", 13
ch.x_axis.title, ch.y_axis.title = "bpp", "PSNR (dB)"
ch.x_axis.delete = ch.y_axis.delete = False
s = Series(Reference(ws, min_col=3, min_row=hs + 1, max_row=hs + 8), Reference(ws, min_col=2, min_row=hs + 1, max_row=hs + 8), title="S1")
s.marker.symbol = "circle"
ch.series.append(s)
ch.height, ch.width = 7, 13
ws.add_chart(ch, f"G{hs}")

r = section(ws, r, "4. Giải mã có khớp từng bit giữa các máy không (T4 và CPU, torch 2.6 và 2.10)")
r = table(ws, r, ["Cách tính", "Số chỗ lệch", "Trên tổng số", "Kết quả"], [
    ("float32", 79, 7864320, "Hỏng: mọi ảnh sẽ giải mã sai"),
    ("bf16", 1114, 7864320, "Hỏng"),
    ("Số nguyên (int8)", 0, 7864320, "Khớp hoàn toàn trên cả 10 cặp máy"),
], fmt={2: "#,##0", 3: "#,##0"})
ws.cell(row=r - 2, column=4).fill = GOOD
ws.cell(row=r - 3, column=4).fill = BAD
ws.cell(row=r - 4, column=4).fill = BAD
r = note(ws, r - 1, "Kết luận: phần quyết định xác suất (entropy model) phải tính bằng số nguyên. Chưa đo trên L4 và P100.", 6)

r = section(ws, r, "5. Tốc độ giải mã DC-AE (1 ảnh 2048×1360, GPU T4)")
sp = r
r = table(ws, sp, ["Cách chạy", "Giây / ảnh", "Mục tiêu (giây / ảnh)", "Chậm hơn mục tiêu (lần)"], [
    ("fp32", 6.74, 0.5, f"=B{sp + 1}/C{sp + 1}"),
    ("fp16", 4.59, 0.5, f"=B{sp + 2}/C{sp + 2}"),
], fmt={2: "0.00", 3: "0.00", 4: "0.0\"×\""})
for i in (sp + 1, sp + 2):
    ws.cell(row=i, column=2).font = blue
    ws.cell(row=i, column=3).font = blue
ws.cell(row=sp + 1, column=3).comment = Comment("Mục tiêu 20 giây cho 30 ảnh, cả pipeline → khoảng 0.5 giây mỗi ảnh (giả định của em).", "Khoi")
note(ws, r - 1, "fp16 chỉ nhanh hơn 1.5 lần → chỗ chậm có lẽ là vòng lặp Python khi ghép các mảnh ảnh, không phải phép tính chính. Việc đầu tiên: đo xem chậm ở đâu.", 6)

# 4. Approach ---------------------------------------------------------------------
ws = sheet("Cách làm", "Các phần của hệ thống", "Ảnh sơ đồ chi tiết: paper/figures/drawio/1_kien_truc.drawio và 2_pipeline_ma_hoa_giai_ma.drawio", [26, 50, 46, 22])
r = table(ws, 4, ["Phần", "Làm gì", "Vì sao", "Trạng thái"], [
    ("Encoder (máy của em)", "Chọn cách nén cho từng ảnh, chia bit giữa các ảnh, thử nhiều seed và chọn kết quả đẹp nhất", "Encoder không bị khóa, sửa được cả sau 01/03", "Chưa làm"),
    ("Mã hóa xác suất bằng số nguyên", "Đọc/ghi bitstream bằng phép tính số nguyên trên CPU", "Để máy chủ giải mã ra đúng từng bit", "Đã có (rANS + mạng int8)"),
    ("Codec S1", "Nén ảnh ở dạng đặc trưng (latent) của DC-AE, 8 mức nén", "Phần nén chính, chiếm ít bit", "Đã xong (200k bước)"),
    ("Bridge (tạo ảnh 1 bước)", "Từ latent mờ, tạo ra latent sắc nét bằng mô hình SANA-0.6B", "Làm ảnh đẹp, giống ảnh thật", "Đang làm (cho CVPR)"),
    ("DC-AE decoder", "Chuyển latent thành ảnh", "Có sẵn, đã viết lại bằng torch thuần", "Đã có, nhưng còn chậm"),
    ("Lớp bù sai số (residual)", "Mã hóa phần chênh lệch giữa ảnh gốc và ảnh AI", "Tăng độ giống gốc ở 0.15 và 0.3 bpp", "Làm sau 16/11"),
    ("Mode VVC-SCC", "Nén ảnh màn hình bằng codec có sẵn trong devkit", "Chữ rõ hơn, PSNR cao hơn 1–3 dB", "Chỉ cần ghép vào"),
    ("Thước đo thay ban giám khảo (Q̂)", "Dự đoán người chấm thích ảnh nào, học từ dữ liệu chấm của CLIC 2021–2024", "Encoder cần một thước đo để chọn", "Chưa làm"),
])
for i in range(5, r - 1):
    st = ws.cell(row=i, column=4)
    st.fill = GOOD if st.value.startswith(("Đã", "Chỉ")) else WARN

# 5. Roadmap ----------------------------------------------------------------------
ws = sheet("Lộ trình", "Lộ trình", "Ngày dạng dd/mm/yyyy. Cột \"Số ngày\" là công thức.", [40, 13, 13, 10, 40])
plan = [("Làm phần chính cho bài CVPR (S1 + bridge)", date(2026, 10, 3), date(2026, 11, 16), "Mốc: đo lại cách AI + bù với phần AI thật"),
        ("Khung nộp bài, nộp thử lần 1–2", date(2026, 11, 17), date(2026, 12, 1), "Kiểm tra cách máy chủ tính dung lượng"),
        ("Mã hóa số nguyên, nộp thử lần 3", date(2026, 12, 1), date(2026, 12, 20), "Mốc: PSNR máy chủ = PSNR ở máy em"),
        ("Lớp bù sai số + chia bit cho cả bộ ảnh", date(2026, 12, 15), date(2027, 1, 10), ""),
        ("Tăng tốc, giảm dung lượng, mức 0.15/0.3", date(2027, 1, 5), date(2027, 1, 31), "Mục tiêu: 30 ảnh ≤ 20 giây"),
        ("Mẹo phía encoder, 3 phiên bản, tự chấm", date(2027, 2, 1), date(2027, 2, 22), ""),
        ("Khóa decoder", date(2027, 2, 22), date(2027, 2, 22), "Chừa 1 tuần dự phòng"),
        ("Hạn nộp decoder (validation)", date(2027, 3, 1), date(2027, 3, 1), "Hạn chính"),
        ("Nộp kết quả trên bộ ảnh test", date(2027, 3, 9), date(2027, 3, 9), "Chỉ chạy encoder")]
rows = [(t, s, e, f"=C{5 + i}-B{5 + i}+1", n) for i, (t, s, e, n) in enumerate(plan)]
r = table(ws, 4, ["Việc", "Bắt đầu", "Kết thúc", "Số ngày", "Ghi chú"], rows, fmt={2: "dd/mm/yyyy", 3: "dd/mm/yyyy", 4: "0"})
for c in range(1, 6):
    ws.cell(row=12, column=c).font = BOLD
    ws.cell(row=12, column=c).fill = WARN

# 6. Risks ------------------------------------------------------------------------
ws = sheet("Rủi ro", "Rủi ro và cách xử lý", "Mức: Cao / Trung bình", [42, 13, 46, 38])
r = table(ws, 4, ["Rủi ro", "Mức", "Cách xử lý", "Nếu vẫn không được"], [
    ("Decoder chậm hơn mục tiêu 9–13 lần", "Cao", "Tìm chỗ chậm, sửa phần ghép ảnh, dùng bf16 trên L4", "Làm decoder nhỏ hơn"),
    ("Phần AI thật ở 0.03 bpp kém hơn dự tính", "Cao", "Đo lại ngay khi có bridge", "Dành nhiều bit hơn cho lớp bù"),
    ("Máy chủ giải mã sai vì phép tính float", "Trung bình", "Dùng số nguyên; nộp thử sớm để so PSNR", ""),
    ("Ở 0.3 bpp thua VTM về độ giống gốc", "Trung bình", "Lớp bù sai số học bằng mạng nơ-ron", "Một phiên bản thiên về giống gốc"),
    ("Khuôn mặt bị đổi người", "Trung bình", "Phạt sai khác trong thước đo Q̂", "Encoder chọn cách ít \"bịa\" hơn cho ảnh đó"),
    ("Thước đo Q̂ không giống ý người chấm", "Trung bình", "Học từ dữ liệu chấm của CLIC; tự chấm ~300 cặp", "3 phiên bản để chia rủi ro"),
    ("Thiếu GPU vì trùng lịch CVPR", "Trung bình", "Phần CLIC chủ yếu là fine-tune, chạy được trên Kaggle", "Thuê L4 theo giờ"),
])
for i in range(5, r - 1):
    c = ws.cell(row=i, column=2)
    c.fill = BAD if c.value == "Cao" else WARN

# 7. Questions --------------------------------------------------------------------
ws = sheet("Câu hỏi", "Câu hỏi xin góp ý", "Ô màu vàng: anh/chị ghi ý kiến vào đây. Dòng 5 là ví dụ cách ghi.", [5, 70, 50])
qs = [
    "Lớp bù sai số: nên mã hóa phần chênh lệch (ảnh gốc − ảnh AI), hay mã hóa thẳng ảnh gốc và dùng ảnh AI làm gợi ý? Cách sau thường tốt hơn nhưng giải mã chậm hơn.",
    "Lớp bù nên học bằng MSE + LPIPS, hay cần thêm GAN? Có kinh nghiệm gì về việc lớp bù làm ảnh mờ lại không?",
    "Thước đo Q̂: học từ dữ liệu chấm CLIC 2021–2024 có ổn không, khi các năm đó gần như không có codec tạo ảnh bằng AI?",
    "Với ~0.5 giây mỗi ảnh, nên cố làm bridge 1 bước, hay chấp nhận 2–4 bước và tăng tốc DC-AE?",
    "Em làm một mình và CVPR chiếm tháng 10–11. Nếu phải cắt, em định cắt phần thử seed/tối ưu latent trước. Thứ tự này ổn không?",
    "3 phiên bản chia theo \"giống gốc ↔ đẹp\" có phải cách chia rủi ro tốt nhất không?",
]
r = table(ws, 4, ["#", "Câu hỏi", "Ý kiến của anh/chị"], [(0, "(Ví dụ)", "Nên thử cách 2 trước, vì ...")] + [(i, q, None) for i, q in enumerate(qs, 1)])
for c in (2, 3):
    ws.cell(row=5, column=c).font = SUB
for i in range(6, 6 + len(qs)):
    ws.cell(row=i, column=3).fill = INPUT_FILL
    ws.row_dimensions[i].height = 45

for w in wb.worksheets:
    w.sheet_properties.pageSetUpPr.fitToPage = True
wb.save(OUT)
print("saved", OUT)
